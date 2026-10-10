"""Build the column-ingestion Excel templates from their YAML description.

The YAML holds the content (tabs, fields, notes, examples, dropdowns); this module
holds the layout and styling, so the two can change independently. One description
builds several templates: anything marked `only: [<template>, ...]` appears only in
those templates.
"""

import shutil
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BOLD = Font(bold=True)
LINK = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center")
_THIN = Side(style="thin", color="999999")
BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

EXAMPLE_FILL = "F2F2F2"
MEANING_FILL = "FFFFFF"
FORMAT_FILL = "F7F7F7"
GUIDE_TEXT = "404040"
WARNING_TEXT = "C00000"
LABEL_TEXT = "7F7F7F"
# Header columns narrower than this are widened so the header text fits.
MIN_WIDTH = 12
# Paragraphs on the README tab are given a taller row to fit wrapped text.
PARAGRAPH_HEIGHT = 32
# A step longer than this wraps onto a second line.
STEP_WRAP_LENGTH = 95
# Checklist boxes: a dropdown, since openpyxl cannot write Excel's native checkboxes.
UNCHECKED, CHECKED = "☐", "☑"
DONE_TEXT = "A6A6A6"


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _note(text: str) -> Comment:
    comment = Comment(text, "Macrostrat")
    comment.width = 320
    comment.height = 150
    return comment


def load_spec(path) -> dict:
    """Read a template description from YAML."""
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


#: Appended to the published template's name for a freshly built copy.
DRAFT_SUFFIX = "-new"


def template_names(spec_path) -> list[str]:
    """The templates the spec builds, in order; the first is the default."""
    return [t["name"] for t in load_spec(spec_path)["templates"]]


def _template(spec: dict, name: str | None) -> dict:
    if name is None:
        return spec["templates"][0]
    for template in spec["templates"]:
        if template["name"] == name:
            return template
    raise ValueError(f"No template named {name!r}")


def published_path(spec_path, template: str | None = None) -> Path:
    """A published template: its `output`, next to the YAML file."""
    spec_path = Path(spec_path)
    return spec_path.parent / _template(load_spec(spec_path), template)["output"]


def draft_path(spec_path, template: str | None = None) -> Path:
    """Where a new build goes by default, so the published template is never overwritten."""
    published = published_path(spec_path, template)
    return published.with_name(f"{published.stem}{DRAFT_SUFFIX}{published.suffix}")


def build_template(spec_path, output=None, template: str | None = None) -> Path:
    """Build one template described by `spec_path` and save it.

    Writes to `output`, or by default to a `-new` copy beside the published template
    (`draft_path`). Returns the path written.
    """
    spec_path = Path(spec_path)
    spec = load_spec(spec_path)
    name = _template(spec, template)["name"]
    workbook = TemplateBuilder(spec, name).build()
    out = Path(output) if output is not None else draft_path(spec_path, name)
    workbook.save(out)
    return out


def build_templates(spec_path) -> list[Path]:
    """Build every template to its `-new` copy."""
    return [build_template(spec_path, template=n) for n in template_names(spec_path)]


def promote_templates(spec_path) -> list[Path]:
    """Replace each published template with its `-new` copy, which is moved, not copied."""
    promoted = []
    for name in template_names(spec_path):
        draft = draft_path(spec_path, name)
        if not draft.exists():
            raise FileNotFoundError(f"No new build of {name!r}: build the templates first")
        published = published_path(spec_path, name)
        shutil.move(draft, published)
        promoted.append(published)
    return promoted


def included(item, template: str) -> bool:
    """Whether a field, row, tab or block belongs in `template`."""
    if not isinstance(item, dict) or "only" not in item:
        return True
    return template in item["only"]


class TemplateBuilder:
    """Turns a loaded template description into an openpyxl `Workbook` for one template."""

    def __init__(self, spec: dict, template: str | None = None):
        self.spec = spec
        self.template = _template(spec, template)
        self.kinds = spec["kinds"]
        self.max_rows = spec["max_rows"]
        self.example_text = spec["example_tabs"]
        self.wb = Workbook()

    def _only(self, items: list) -> list:
        return [item for item in items if included(item, self.template["name"])]

    def build(self) -> Workbook:
        self._readme(self.spec["readme"])
        for sheet in self._only(self.spec["sheets"]):
            if sheet.get("layout") == "key_value":
                self._key_value_sheet(sheet)
            else:
                self._field_sheet(sheet)
        return self.wb

    # ------------------------------------------------------------ README

    def _readme(self, page: dict):
        ws = self.wb.active
        ws.title = page["name"]
        ws.sheet_properties.tabColor = page["tab_color"]
        for column, width in (("A", 6), ("B", 30), ("C", 95)):
            ws.column_dimensions[column].width = width

        ws.cell(1, 2, self.template["title"]).font = Font(bold=True, size=16)
        row = 2
        for block in self._only(page["blocks"]):
            row = self._readme_block(ws, row, block)

    def _readme_block(self, ws, row: int, block: dict) -> int:
        """Write one README block at `row`; return the next free row."""
        if "link" in block:
            cell = ws.cell(row, 2, block["link"])
            cell.hyperlink = block["url"]
            cell.font = LINK
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            return row + 1
        if "heading" in block:
            row += 1
            ws.cell(row, 2, block["heading"]).font = Font(bold=True, size=12)
            return row + 1
        if "paragraph" in block:
            ws.cell(row, 2, block["paragraph"]).alignment = WRAP
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            ws.row_dimensions[row].height = PARAGRAPH_HEIGHT
            return row + 1
        if "legend" in block:
            for name in block["legend"]:
                kind = self.kinds[name]
                cell = ws.cell(row, 2, kind["label"])
                cell.fill = _fill(kind["color"])
                cell.font = BOLD
                cell.border = BOX
                ws.cell(row, 3, kind["description"])
                row += 1
            return row
        if "steps" in block:
            for number, sheet, text in self._entries(block["steps"]):
                ws.cell(row, 1, str(number)).font = BOLD
                ws.cell(row, 2, sheet).font = BOLD
                ws.cell(row, 3, text).alignment = WRAP
                height = PARAGRAPH_HEIGHT if len(text) > STEP_WRAP_LENGTH else None
                ws.row_dimensions[row].height = height
                row += 1
            return row
        if "checklist" in block:
            return self._checklist(ws, row, block["checklist"])
        raise ValueError(f"Unknown README block: {block}")

    def _checklist(self, ws, row: int, groups: list) -> int:
        """Checklist items grouped by tab, each with a box ticked from a dropdown."""
        first, boxes = row, []
        for group in self._only(groups):
            ws.cell(row, 2, group["tab"]).font = Font(bold=True, italic=True)
            row += 1
            for item in self._entries(group["items"]):
                box = ws.cell(row, 1, UNCHECKED)
                box.alignment = CENTER
                box.border = BOX
                boxes.append(box.coordinate)
                ws.cell(row, 2, item)
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
                row += 1
        validation = DataValidation(
            type="list", formula1=f'"{UNCHECKED},{CHECKED}"', allow_blank=True
        )
        for box in boxes:
            validation.add(box)
        ws.add_data_validation(validation)
        done = FormulaRule(
            formula=[f'$A{first}="{CHECKED}"'],
            fill=_fill("F2F2F2"),
            font=Font(color=DONE_TEXT, strike=True),
        )
        ws.conditional_formatting.add(f"B{first}:C{row - 1}", done)
        return row

    def _entries(self, entries: list) -> list:
        """A step or checklist list, where an entry may be `{item: ..., only: [...]}`."""
        return [e["item"] if isinstance(e, dict) else e for e in self._only(entries)]

    # ------------------------------------------------------------ field tabs

    def _header(self, ws, fields: list[dict], tab_color: str, first_column=1):
        ws.sheet_properties.tabColor = tab_color
        for i, field in enumerate(fields, first_column):
            cell = ws.cell(1, i, field["name"])
            cell.font = BOLD
            cell.fill = _fill(self.kinds[field["kind"]]["color"])
            cell.border = BOX
            cell.alignment = CENTER
            cell.comment = _note(_note_text(field))
            width = max(field["width"], MIN_WIDTH)
            ws.column_dimensions[cell.column_letter].width = width

    def _field_sheet(self, sheet: dict):
        fields = self._only(sheet["fields"])
        ws = self.wb.create_sheet(sheet["name"])
        self._header(ws, fields, sheet["tab_color"])
        ws.freeze_panes = "A2"

        for i, field in enumerate(fields, 1):
            if "validation" in field:
                self._validate(ws, get_column_letter(i), field["validation"])
        if "highlight_missing" in sheet:
            self._highlight_missing(ws, fields, sheet["highlight_missing"])
        if "example" in sheet:
            self._field_example(sheet["name"], fields, sheet["example"])

    def _validate(self, ws, column: str, rule: dict):
        kind = rule["type"]
        options = {}
        if kind == "list":
            formula1 = f"={rule['range']}"
        else:
            formula1 = str(rule["value"])
            options["operator"] = rule["operator"]
            if "max" in rule:
                options["formula2"] = str(rule["max"])
        # A non-strict rule warns about an unlisted value but still accepts it.
        error_style = "stop"
        if rule.get("strict") is False:
            error_style = "warning"
        validation = DataValidation(
            type=kind,
            formula1=formula1,
            allow_blank=True,
            showErrorMessage=True,
            errorStyle=error_style,
            prompt=rule["prompt"],
            showInputMessage=True,
            **options,
        )
        validation.add(f"{column}2:{column}{self.max_rows}")
        ws.add_data_validation(validation)

    def _highlight_missing(self, ws, fields: list[dict], rule: dict):
        letters = {f["name"]: get_column_letter(i) for i, f in enumerate(fields, 1)}
        targets = sorted(letters[name] for name in rule["fields"])
        trigger = letters[rule["when_filled"]]
        blanks = ",".join(f'${letter}2=""' for letter in targets)
        formula = f'AND(${trigger}2<>"",{blanks})'
        cells = f"{targets[0]}2:{targets[-1]}{self.max_rows}"
        ws.conditional_formatting.add(
            cells, FormulaRule(formula=[formula], fill=_fill(rule["color"]))
        )

    def _field_example(self, name: str, fields: list[dict], example: dict):
        text = self.example_text
        ws = self.wb.create_sheet(f"{name} (example)")
        ws.sheet_properties.tabColor = example["tab_color"]
        ws.column_dimensions["A"].width = 13
        ws.cell(1, 1, "").border = BOX
        self._header(ws, fields, example["tab_color"], first_column=2)

        guide_rows = ((text["meaning_label"], "meaning"), (text["format_label"], "format"))
        for row, (label, key) in enumerate(guide_rows, 2):
            label_cell = ws.cell(row, 1, label)
            label_cell.font = BOLD
            label_cell.alignment = WRAP
            is_meaning = key == "meaning"
            background = FORMAT_FILL
            if is_meaning:
                background = MEANING_FILL
            for i, field in enumerate(fields, 2):
                cell = ws.cell(row, i, field[key])
                cell.alignment = WRAP
                cell.border = BOX
                cell.font = Font(italic=is_meaning, size=9, color=GUIDE_TEXT)
                cell.fill = _fill(background)
                if key == "format" and "link" in field:
                    cell.hyperlink = field["link"]
                    cell.font = Font(size=9, color=LINK.color, underline="single")
        ws.row_dimensions[2].height = 95
        ws.row_dimensions[3].height = 30

        rows = self._only(example["rows"])
        for row, values in enumerate(rows, 4):
            ws.cell(row, 1, text["row_label"]).font = Font(bold=True, color=LABEL_TEXT)
            for i, field in enumerate(fields, 2):
                value = values.get(field["name"])
                cell = ws.cell(row, i, value)
                cell.fill = _fill(EXAMPLE_FILL)
                cell.border = BOX
                long_text = isinstance(value, str) and len(value) > 30
                cell.alignment = Alignment(vertical="top", wrap_text=long_text)
        footer = ws.cell(len(rows) + 5, 1, text["footer"])
        footer.font = Font(italic=True, color=WARNING_TEXT)
        ws.freeze_panes = "B4"

    # ------------------------------------------------------------ key/value tab

    def _key_value_sheet(self, sheet: dict):
        rows = self._only(sheet["rows"])
        ws = self.wb.create_sheet(sheet["name"])
        ws.sheet_properties.tabColor = sheet["tab_color"]
        for i, row in enumerate(rows, 1):
            key = ws.cell(i, 1, row["key"])
            key.font = BOLD
            key.comment = _note(_note_text(row))
            key.border = BOX
            value = ws.cell(i, 2, row.get("default"))
            value.fill = _fill(self.kinds[row["kind"]]["color"])
            value.border = BOX
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 60
        for i, row in enumerate(rows, 1):
            if "choices" in row:
                choices = ",".join(row["choices"])
                validation = DataValidation(type="list", formula1=f'"{choices}"')
                validation.add(f"B{i}")
                ws.add_data_validation(validation)
        if "example" in sheet:
            self._key_value_example(sheet["name"], rows, sheet["example"])

    def _key_value_example(self, name: str, rows: list[dict], example: dict):
        ws = self.wb.create_sheet(f"{name} (example)")
        ws.sheet_properties.tabColor = example["tab_color"]
        for j, (title, width) in enumerate(example["columns"], 1):
            cell = ws.cell(1, j, title)
            cell.font = BOLD
            cell.border = BOX
            ws.column_dimensions[cell.column_letter].width = width
        for i, row in enumerate(rows, 2):
            key = ws.cell(i, 1, row["key"])
            key.font = BOLD
            key.fill = _fill(self.kinds[row["kind"]]["color"])
            key.border = BOX
            for j, value in enumerate((row["example"], row["meaning"], row["format"]), 2):
                cell = ws.cell(i, j, value)
                cell.alignment = WRAP
                cell.border = BOX
                if j == 2:
                    cell.fill = _fill(EXAMPLE_FILL)
                else:
                    cell.font = Font(italic=(j == 3), size=9, color=GUIDE_TEXT)
        footer = ws.cell(len(rows) + 3, 1, example["footer"])
        footer.font = Font(italic=True, color=WARNING_TEXT)
        ws.freeze_panes = "A2"


def _note_text(field: dict) -> str:
    text = f"{field['meaning']}\nFormat: {field['format']}"
    if "link" in field:
        text += f"\nSee {field['link']}"
    return text
