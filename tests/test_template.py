"""Each published template must be exactly what its YAML description builds."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from column_template import build_template, published_path, template_names

SPEC = Path(__file__).parent.parent / "Excel Templates" / "column-ingestion-template.yaml"
OUT_OF_DATE = (
    "A published template no longer matches the YAML. Run `make template`, review the "
    "-new copies, then `make promote` to replace the published templates with them."
)


def _style(cell):
    return (
        cell.font.b,
        cell.font.i,
        cell.font.sz,
        cell.font.color and cell.font.color.rgb,
        cell.fill.fill_type,
        cell.fill.fgColor.rgb,
        cell.border.left.style,
        cell.border.left.color and cell.border.left.color.rgb,
        cell.alignment.horizontal,
        cell.alignment.vertical,
        cell.alignment.wrap_text,
    )


def _validations(ws):
    return sorted(
        (
            str(dv.sqref),
            dv.type,
            dv.formula1,
            dv.formula2,
            dv.operator,
            dv.allow_blank,
            dv.showErrorMessage,
            dv.errorStyle,
            dv.prompt,
            dv.showInputMessage,
        )
        for dv in ws.data_validations.dataValidation
    )


def _highlights(ws):
    return sorted(
        (str(cf.sqref), rule.formula[0], rule.dxf.fill.fgColor.rgb)
        for cf in ws.conditional_formatting
        for rule in cf.rules
    )


def describe(path) -> dict:
    """Everything about a workbook that the template defines."""
    wb = load_workbook(path)
    sheets = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for cell in row:
                note = cell.comment and (cell.comment.text, cell.comment.author)
                link = cell.hyperlink and cell.hyperlink.target
                cells[cell.coordinate] = (cell.value, _style(cell), note, link)
        sheets[ws.title] = {
            "cells": cells,
            "widths": {k: d.width for k, d in ws.column_dimensions.items()},
            "heights": {k: d.height for k, d in ws.row_dimensions.items() if d.height},
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "freeze": ws.freeze_panes,
            "tab": ws.sheet_properties.tabColor and ws.sheet_properties.tabColor.rgb,
            "validations": _validations(ws),
            "highlights": _highlights(ws),
        }
    return {"order": wb.sheetnames, "sheets": sheets}


@pytest.mark.parametrize("template", template_names(SPEC))
def test_template_matches_yaml(tmp_path, template):
    built = build_template(SPEC, tmp_path / "built.xlsx", template)
    expected, actual = describe(published_path(SPEC, template)), describe(built)
    assert actual["order"] == expected["order"], OUT_OF_DATE
    for name in expected["order"]:
        assert actual["sheets"][name] == expected["sheets"][name], (
            f"{OUT_OF_DATE} (First difference on the `{name}` tab.)"
        )
